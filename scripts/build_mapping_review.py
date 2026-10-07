"""Build the review workbook for the drafted reference mappings (open items O-P3-2, O-P3-9, O-P4-1, O-P5-4, O-P6-7, O5).

One sheet per mapping: the draft row as loaded in ClickHouse `default`, its usage in the warehouse, a FLAG that puts the
rows needing attention first, and yellow review columns (STATUS, CORRECTED_VALUE, COMMENT) for the owner to fill in.
Read-only on ClickHouse. The workbook is written outside git (static_mappings/review/, *.xlsx is ignored).

Usage:  python scripts/build_mapping_review.py
"""
import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).resolve().parent))
from ch_env import client  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "static_mappings" / "review" / f"mapping_review_{date.today():%Y-%m-%d}.xlsx"

FONT = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
INPUT_FILL = PatternFill("solid", fgColor="FFFF00")
FLAG_FILL = PatternFill("solid", fgColor="FCE4D6")
THIN = Side(style="thin", color="BFBFBF")
REVIEW_COLS = ["STATUS", "CORRECTED_VALUE", "COMMENT"]

# (sheet, open item, owner, what to check, corrected-value meaning, SQL). Each SQL returns FLAG first ('' = nothing to check),
# then the draft columns, then usage columns; rows are sorted flagged first, then by usage.
SHEETS = [
    ("FS accounts (inferred)", "O-P3-2", "Finance",
     "97 accounts mapped by inference from 6-digit sibling groups. Check FS_LINE / FS_CAPTION.",
     "Correct FS_LINE (and caption if different)",
     """
     select if(m.FS_LINE ilike '%VAT%' and n.name ilike '%withholding%', 'Borderline: withholding tax on the VAT payable line', '') as FLAG,
            m.ORACLE_CODE, n.name as ACCOUNT_NAME, n.account_type as ACCOUNT_TYPE, m.FS_TYPE, m.FS_ELEMENT, m.FS_CATEGORY, m.FS_CAPTION, m.FS_LINE,
            ifNull(u.branches, '') as BRANCHES, ifNull(u.lines, 0) as JOURNAL_LINES, round(ifNull(u.gross, 0), 2) as GROSS_SAR
     from default.map_fs_account as m
     left join (select toString(natural_account) as code, any(natural_account_name) as name, any(account_type) as account_type
                from gold.dim_gl_account group by code) as n on n.code = toString(m.ORACLE_CODE)
     left join (select toString(a.natural_account) as code, arrayStringConcat(arraySort(groupUniqArray(toString(j.branch_key))), ' ') as branches,
                       count() as lines, sum(j.debit) + sum(j.credit) as gross
                from gold.fact_gl_journal_line as j inner join gold.dim_gl_account as a on a.gl_account_key = j.gl_account_key
                group by code) as u on u.code = toString(m.ORACLE_CODE)
     where m.MAPPED_IN = 'inferred'
     order by FLAG = '', GROSS_SAR desc
     settings join_use_nulls = 1"""),
    ("FS line order", "O-P3-9", "Finance",
     "Sort order and statement group of each FS type / element / category / caption.",
     "Correct SORT_ORDER or STATEMENT_GROUP",
     """
     select '' as FLAG, LEVEL, VALUE, SORT_ORDER, STATEMENT_GROUP from default.map_fs_line_order order by LEVEL, SORT_ORDER"""),
    ("Budget to FS line", "O-P3-9", "Finance",
     "Which FS level and value each budget line item compares with, and its care type.",
     "Correct MATCH_LEVEL / MATCH_VALUE / CARE_TYPE",
     """
     select if(ifNull(MATCH_VALUE, '') = '', 'No FS match', '') as FLAG, LINE_ITEM_CODE, MATCH_LEVEL, MATCH_VALUE, CARE_TYPE
     from default.map_budget_fs_line order by FLAG = '', LINE_ITEM_CODE"""),
    ("Fusion specialty", "O-P3-9", "BI manager",
     "Unified department of each Fusion specialty (segment 3). Blank = Unknown; flagged where GL amounts use it (placeholders 000000000 and NO VAL excluded).",
     "Unified department",
     """
     select multiIf(ifNull(m.UNIFIED_DEPARTMENT, '') = '' and ifNull(u.gross, 0) > 0
                    and m.SPECIALTY_CODE not in ('000000000', 'NO VAL'), 'Blank, used in GL', '') as FLAG,
            m.SPECIALTY_CODE, m.SPECIALTY_NAME, m.UNIFIED_DEPARTMENT,
            ifNull(u.lines, 0) as JOURNAL_LINES, round(ifNull(u.gross, 0), 2) as GROSS_SAR
     from default.map_fusion_specialty_unified as m
     left join (select toString(a.specialty_code) as code, count() as lines, sum(j.debit) + sum(j.credit) as gross
                from gold.fact_gl_journal_line as j inner join gold.dim_gl_account as a on a.gl_account_key = j.gl_account_key
                group by code) as u on u.code = toString(m.SPECIALTY_CODE)
     order by FLAG = '', GROSS_SAR desc
     settings join_use_nulls = 1"""),
    ("Pay category", "O-P4-1", "HR / payroll",
     "Pay category of each Oasis transaction type (and payable type) and Fusion element. 2026 amounts as loaded.",
     "Correct pay category",
     """
     with oasis as (
         select 'oasis' as SOURCE, trx_type as SOURCE_CODE, payable_type as PAYABLE_TYPE, count() as LINES_2026, sum(amount) as AMOUNT_2026
         from stg.stg_oasis__payroll_transactions where status = 'C' and payroll_month >= 202601 group by trx_type, payable_type
     ),
     fusion as (
         select 'fusion' as SOURCE, e.element_name as SOURCE_CODE, '' as PAYABLE_TYPE, count() as LINES_2026, sum(r.result_value) as AMOUNT_2026
         from stg.stg_fusion__payroll_run_results as r
         inner join (select input_value_id from stg.stg_fusion__payroll_input_values where input_value_base_name = 'Pay Value') as i
             on i.input_value_id = r.input_value_id
         inner join stg.stg_fusion__payroll_elements as e on e.element_type_id = r.element_type_id
         where r.payroll_action_status = 'C' and r.effective_date >= '2026-01-01'
         group by e.element_name
     ),
     used as (select * from oasis union all select * from fusion)
     select multiIf(m.PAY_CATEGORY is null, 'Not in mapping (Unmapped)', m.PAY_CATEGORY = 'Unmapped', 'Mapped to Unmapped', '') as FLAG,
            u.SOURCE as SOURCE, u.SOURCE_CODE as SOURCE_CODE, u.PAYABLE_TYPE as PAYABLE_TYPE, ifNull(m.PAY_CATEGORY, '') as PAY_CATEGORY,
            u.LINES_2026 as LINES_2026, round(u.AMOUNT_2026, 2) as AMOUNT_2026
     from used as u
     left join default.map_pay_category as m
         on m.SOURCE = u.SOURCE and upper(trimBoth(m.SOURCE_CODE)) = upper(trimBoth(u.SOURCE_CODE))
        and (u.SOURCE = 'fusion' or upper(trimBoth(m.PAYABLE_TYPE)) = u.PAYABLE_TYPE)
     order by FLAG = '', abs(AMOUNT_2026) desc
     settings join_use_nulls = 1"""),
    ("Store department", "O-P5-4", "BI manager",
     "Store type and unified department of each Oasis and Fusion store. Flagged: unmapped store type, or a clinic / ward / theatre store without department, with 2026 movements.",
     "Correct STORE_TYPE / UNIFIED_DEPARTMENT",
     """
     select multiIf(ifNull(u.lines, 0) > 0 and ifNull(m.STORE_TYPE, '') in ('', 'Unknown', 'Unmapped'), 'Store type unmapped, used',
                    ifNull(u.lines, 0) > 0 and m.STORE_TYPE in ('Clinic', 'Ward', 'Operating room')
                    and ifNull(m.UNIFIED_DEPARTMENT, '') in ('', 'Unknown', 'Not Mapped'), 'Clinical store without department, used', '') as FLAG,
            m.SOURCE, m.BRANCH_ID, m.STORE_CODE, m.STORE_NAME, m.STORE_TYPE, m.UNIFIED_DEPARTMENT,
            ifNull(u.lines, 0) as MOVEMENT_LINES_2026, round(ifNull(u.cost, 0), 2) as ABS_COST_2026
     from default.map_store_department as m
     left join (select s.source_system as src, s.branch_key as br, s.store_code as code, count() as lines, sum(abs(f.cost_amount)) as cost
                from gold.fact_stock_movement as f inner join gold.dim_store as s on s.store_key = f.store_key
                where f.date_key >= 20260101 group by src, br, code) as u
         on u.src = m.SOURCE and u.br = m.BRANCH_ID and u.code = m.STORE_CODE
     order by FLAG = '', ABS_COST_2026 desc
     settings join_use_nulls = 1"""),
    ("Item group", "O-P5-4", "BI manager",
     "Item group of each Fusion HNH Catalog category code. Flagged: blank or Other on a category with items.",
     "Correct ITEM_GROUP",
     """
     select if(ifNull(m.ITEM_GROUP, '') in ('', 'Other', 'Unknown') and ifNull(u.items, 0) > 0, 'Blank / Other / Unknown, used', '') as FLAG,
            m.CATEGORY_CODE, m.ITEM_GROUP, ifNull(u.items, 0) as ITEMS, ifNull(u.example, '') as EXAMPLE_ITEM
     from default.map_item_group as m
     left join (select category_code, count() as items, any(item_description) as example from gold.dim_item group by category_code) as u
         on u.category_code = m.CATEGORY_CODE
     order by FLAG = '', ITEMS desc
     settings join_use_nulls = 1"""),
    ("PG question role", "O-P6-7", "Patient experience",
     "Role of each Press Ganey question per service: Hospital NPS, Physician NPS, or the background attribute it feeds.",
     "Correct ROLE",
     """
     select '' as FLAG, m.SERVICE, m.QUESTION_CODE, ifNull(q.question_en, '') as QUESTION, m.ROLE, ifNull(a.answers, 0) as ANSWERS
     from default.map_pg_question_role as m
     left join (select service_code, question_code, any(question_en) as question_en, any(question_key) as question_key
                from gold.dim_survey_question group by service_code, question_code) as q
         on q.service_code = m.SERVICE and q.question_code = m.QUESTION_CODE
     left join (select question_key, count() as answers from gold.fact_survey_answer group by question_key) as a on a.question_key = q.question_key
     order by m.SERVICE, m.ROLE
     settings join_use_nulls = 1"""),
    ("PG background value", "O-P6-7", "Patient experience",
     "Conformed value of each background answer code.",
     "Correct CONFORMED_VALUE",
     """
     select '' as FLAG, m.SERVICE, m.QUESTION_CODE, ifNull(q.question_en, '') as QUESTION, m.ANSWER_CODE, ifNull(a.label, '') as SOURCE_LABEL,
            m.CONFORMED_VALUE, ifNull(a.answers, 0) as ANSWERS
     from default.map_pg_background_value as m
     left join (select service_code, question_code, any(question_en) as question_en, any(question_key) as question_key
                from gold.dim_survey_question group by service_code, question_code) as q
         on q.service_code = m.SERVICE and q.question_code = m.QUESTION_CODE
     left join (select question_key, answer_code, any(answer_label_en) as label, count() as answers
                from gold.fact_survey_answer group by question_key, answer_code) as a
         on a.question_key = q.question_key and a.answer_code = m.ANSWER_CODE
     order by m.SERVICE, m.QUESTION_CODE, m.ANSWER_CODE
     settings join_use_nulls = 1"""),
    ("Beds not classified", "O5", "BI manager / nursing",
     "Beds with no row in map_bed_classification; they report Not Mapped and never count as Critical (ICU).",
     "Classification (e.g. Critical, Ward, ...)",
     """
     select 'Not classified' as FLAG, b.branch_key as BRANCH_ID, b.bed_location as BED, b.current_ward as WARD, b.bed_class as BED_CLASS,
            b.bed_gender as BED_GENDER, b.current_slot_status as CURRENT_STATUS
     from gold.dim_bed as b
     where b.classification = 'Not Mapped'
     order by BRANCH_ID, WARD, BED"""),
]


def style_header(ws, ncols, review_from):
    for col in range(1, ncols + 1):
        cell = ws.cell(row=1, column=col)
        is_review = col >= review_from
        cell.font = Font(name=FONT, bold=True, color="000000" if is_review else "FFFFFF")
        cell.fill = INPUT_FILL if is_review else HEAD_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "B2"


def write_sheet(wb, c, spec):
    title, item, owner, check, corrected, sql = spec
    result = c.query(sql)
    cols = list(result.column_names)
    rows = result.result_rows
    ws = wb.create_sheet(title[:31])
    headers = cols + REVIEW_COLS
    ws.append(headers)
    for r in rows:
        ws.append(list(r) + [None, None, None])
    review_from = len(cols) + 1
    style_header(ws, len(headers), review_from)
    ws.cell(row=1, column=review_from + 1).comment = Comment(f"Fill only when STATUS = Change: {corrected}.", "review")
    dv = DataValidation(type="list", formula1='"OK,Change,Question"', allow_blank=True)
    ws.add_data_validation(dv)
    last = len(rows) + 1
    if last >= 2:
        dv.add(f"{get_column_letter(review_from)}2:{get_column_letter(review_from)}{last}")
    for i, r in enumerate(rows, start=2):
        flagged = bool(r[0])
        for col in range(1, len(headers) + 1):
            cell = ws.cell(row=i, column=col)
            cell.font = Font(name=FONT, size=10)
            cell.border = Border(bottom=THIN)
            if col >= review_from:
                cell.fill = INPUT_FILL
            elif flagged:
                cell.fill = FLAG_FILL
            if isinstance(cell.value, float):
                cell.number_format = "#,##0.00;(#,##0.00);-"
            elif isinstance(cell.value, int) and cols[col - 1] not in ("BRANCH_ID", "SORT_ORDER", "ORACLE_CODE"):
                cell.number_format = "#,##0;(#,##0);-"
    for col, name in enumerate(headers, start=1):
        width = max([len(str(name))] + [len(str(ws.cell(row=i, column=col).value or "")) for i in range(2, min(last, 300) + 1)])
        ws.column_dimensions[get_column_letter(col)].width = min(max(width + 2, 10), 50)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(last, 1)}"
    flagged_n = sum(1 for r in rows if r[0])
    return title[:31], item, owner, check, len(rows), flagged_n, get_column_letter(review_from)


def main():
    c = client()
    wb = Workbook()
    readme = wb.active
    readme.title = "README"
    summary = [write_sheet(wb, c, s) for s in SHEETS]

    readme["A1"] = "HNH reference mapping review"
    readme["A1"].font = Font(name=FONT, bold=True, size=14)
    readme["A2"] = (f"Built {date.today():%Y-%m-%d} from the draft tables in ClickHouse `default` (read-only). Usage columns show how much "
                    "each row matters in the warehouse. Rows needing attention have a FLAG and are listed first (orange).")
    readme["A4"] = "How to review"
    readme["A4"].font = Font(name=FONT, bold=True)
    steps = [
        "Edit only the yellow columns: STATUS (OK / Change / Question), CORRECTED_VALUE (only when STATUS = Change), COMMENT.",
        "Do not edit, sort away or delete the other columns: they identify the row when the corrections are loaded back.",
        "Return the file; the corrected rows are reloaded into the mapping tables and the models rebuilt.",
        "Example (sheet 'Pay category'): STATUS = Change, CORRECTED_VALUE = Overtime, COMMENT = Night shift premium is overtime pay.",
    ]
    for i, s in enumerate(steps, start=5):
        readme.cell(row=i, column=1, value=f"{i - 4}. {s}")
    hdr_row = 11
    heads = ["Sheet", "Open item", "Owner", "What to check", "Rows", "Flagged", "Reviewed (STATUS filled)", "Changes requested"]
    for col, h in enumerate(heads, start=1):
        cell = readme.cell(row=hdr_row, column=col, value=h)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = HEAD_FILL
    for i, (title, item, owner, check, n, flagged, status_col) in enumerate(summary, start=hdr_row + 1):
        ref = f"'{title}'!{status_col}:{status_col}"
        values = [title, item, owner, check, n, flagged, f'=COUNTA({ref})-1', f'=COUNTIF({ref},"Change")']
        for col, v in enumerate(values, start=1):
            cell = readme.cell(row=i, column=col, value=v)
            cell.font = Font(name=FONT, size=10)
            cell.alignment = Alignment(wrap_text=col == 4, vertical="top")
    for col, w in zip("ABCDEFGH", (24, 10, 18, 70, 8, 9, 14, 12)):
        readme.column_dimensions[col].width = w
    for r in range(2, hdr_row):
        if r != 4:
            readme.cell(row=r, column=1).font = Font(name=FONT, size=10)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(OUT)
    for s in summary:
        print(f"  {s[0]}: {s[4]} rows, {s[5]} flagged")


if __name__ == "__main__":
    sys.exit(main())

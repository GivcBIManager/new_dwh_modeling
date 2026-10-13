"""Build the Excel comparison of the bsc views, old server against new, from the CSV written by
scripts/compare_bsc.ps1 -CsvPath (every key x column with old value, new value, data source and status).

  python scripts/bsc_comparison_xlsx.py <comparison.csv> <out.xlsx>

Sheets: Summary (counts per view, COUNTIFS over Data), Headline KPIs (old and new totals, SUMIFS over Data), Data (one
row per key x column, with difference and % as formulas, and the cause) and Causes (the legend). Causes follow
docs/reconciliation_bsc.md and the spec docs/superpowers/specs/2026-10-12-hnh-bsc-views-design.md.
"""
import csv
import re
import sys
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

FONT = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
DIFF_FILL = PatternFill("solid", fgColor="FCE4D6")
NUM = '#,##0.00;-#,##0.00;"-"'
INT = '#,##0;-#,##0;"-"'
PCT = '0.0%;-0.0%;"-"'

# cause code -> explanation (Causes sheet)
CAUSES = {
    "F7": "The old view summed only GL periods still open (GL_STATUS = 'O'); the new one rolls the fiscal year from the brought-forward rows. Checked against the old raw GL: the new balances are right.",
    "F6": "The old operating-expense filter missed depreciation ('depreciation On Rou' against lower-case text, misspelt 'deprecition'); the new one excludes the whole depreciation category.",
    "B5": "From a branch's first Fusion GL month the new view takes Fusion only; the old view kept the Oasis feed (revenue without costs). Fusion holds almost nothing for September 2026 yet (O-B9).",
    "O-B10": "Khamis has two near-identical Fusion opening batches for Aug-26 (248206 and 254343), so its balance sheet is doubled from September 2026. Finance to confirm which batch is real.",
    "BS-SEP26": "Balance sheet at 2026-09: Oasis branches are higher because the old view summed only open GL periods (F7); Khamis is doubled by its duplicate Fusion opening batch (O-B10); receivables step up at each branch's cutover (O-B10).",
    "CONS": "Consumables budget: the Consumables line only (2026-10-13); the old view also added Cost Of Medicines.",
    "BUDGET": "Months with a budget but no actuals: the old view dropped them (ANY LEFT JOIN).",
    "GL-KEYS": "Old-only GL keys: padded or double-spaced labels the new mapping normalises, zero rows, and Oasis postings after a branch's cutover.",
    "GL-NEW": "New-only GL keys: Fusion accounts and vocabulary, and opening-only rows the old view did not list.",
    "GL-FWD": "forwarded_balance is now the month's brought-forward (Oasis) or go-live (Fusion) amount, not a running total; OpeningBalance + forwarded_balance keeps the old meaning.",
    "GL-OPEN": "The old opening-balance view covered open GL periods only; the new one covers every month from 2022.",
    "REV": "Revenue: SSAS Revenue on corrected charge lines (no post-invoice discount rows); payer splits use the episode payer as before.",
    "ADM": "Countable admissions (over one hour, not a wrong admission); the old rule also dropped stays whose last bed was in an excluded ward. Discharges, LOS, readmissions and ICU follow.",
    "SLOTS": "legacy_capacity_slots counts a doctor-day per department, so a doctor in two clinics on one day counts twice (O-B3, BI manager).",
    "EPIS": "True OP episodes (SSAS Episodes); the old view counted non-follow-up OP visits.",
    "VISITS": "Visits: cancellation by outcome group, guarded wait minutes, episode referral policy, booked doctor for high value (SSAS definitions).",
    "TARGET": "Targets: gold fact_target_daily (most likely); differences come from budget rows added or changed since the old view was built.",
    "BEDS-ALOS": "ALOS: the month's discharges (SSAS ALOS); the old view averaged the to-date stay of every stay open in the month.",
    "BEDS": "Beds: countable stays; average available inpatient beds of the month; bed nights from daily bed occupancy of inpatient, non-excluded wards.",
    "RCM-CLAIMS": "Claims: latest submission, cancelled claims excluded; not_submitted is now the amount pending adjudication; approved and rejected are outcome-aware, and claims without an insurer response are no longer approved (O-P2B-9: Abha and Muhayil pull gaps).",
    "RCM-PREAUTH": "Pre-authorisation: every service line except payer advances (unsent, error and pended included); approved = approved with a final response; turnaround = sent to final response, capped at 1,440 minutes (old 480).",
    "WF-2025": "Before 2026: paid headcount from gold payroll (gross pay > 0, no parallel run) plus revenue-only doctors; the old rule was payroll or revenue per staff member (within about 2%).",
    "WF-2026": "From 2026: Fusion month-end headcount (the old view broke after each branch's payroll cutover); doctors in the headcount are not part-time; Fusion leavers for branches 6-8.",
    "CITY": "Khamis's city is Aseer in gold (old: Asser), as decided.",
    "SAME": "Equal within 0.5%.",
}

FIN_BS = {"assets", "assets_opening_balance", "account_receivable_actual", "account_receivable_actual_opening_balance"}
CUSTOMER = {
    "REV": {"revenue", "op_revenue", "non_ltc_ip_revenue", "ancillary_services_revenue", "non_moh_revenue",
            "insurance_revenue", "referrals_revenue", "credit_revenue", "high_value_service_revenue"},
    "ADM": {"admissions", "discharges", "alos_non_ltc", "readmission_within_30days", "readmitted_to_ICU_with_48hrs",
            "insurance_admissions", "icu_discharge"},
    "SLOTS": {"available_slots"},
    "EPIS": {"op_episodes"},
}


def cause(view, column, source, status, branch, period):
    if status == "equal":
        return "SAME"
    if view == "vw_hospital":
        return "CITY"
    if view == "vw_financial":
        if status == "only new":
            return "BUDGET"
        if column == "consumables_budget":
            return "CONS"
        if source == "fusion":
            if branch == "2" and column in FIN_BS:
                return "O-B10"
            return "B5"
        if column in FIN_BS:
            return "F7"
        return "F6"
    if view in ("vw_gl_transactions", "vw_gl_openning_balance"):
        if status == "only old":
            return "GL-KEYS"
        if status == "only new":
            return "GL-OPEN" if view == "vw_gl_openning_balance" else "GL-NEW"
        if column == "forwarded_balance":
            return "GL-FWD"
        if source == "fusion":
            return "O-B10" if branch == "2" else "B5"
        return "F7"
    if view == "vw_customer":
        for code, cols in CUSTOMER.items():
            if column in cols:
                return code
        return "TARGET" if column.endswith("_target") else "VISITS"
    if view == "vw_hospital_beds_utilization":
        return "BEDS-ALOS" if "alos" in column else "BEDS"
    if view == "vw_rcm":
        return "RCM-PREAUTH" if column in {"total_orders", "approved_orders", "approved_not_done",
                                            "approved_not_done_pct", "avg_tat_minutes"} else "RCM-CLAIMS"
    if view == "vw_work_force":
        return "WF-2026" if source == "both" or period >= "2026" else "WF-2025"
    return ""


def split_key(view, key):
    parts = key.split("|")
    period = next((p for p in parts if re.fullmatch(r"\d{4}-\d{2}-\d{2}", p)), "")
    if not period and view in ("targets", "patient_satisfaction") and len(parts) > 1 and re.fullmatch(r"\d{6}", parts[1]):
        period = f"{parts[1][:4]}-{parts[1][4:]}-01"
    branch = parts[0] if view != "vw_date_dim" else ""
    return branch, period


def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)


def main():
    src, out = sys.argv[1], sys.argv[2]
    with open(src, encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.DictReader(fh))

    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    headline = wb.create_sheet("Headline KPIs")
    data = wb.create_sheet("Data")
    legend = wb.create_sheet("Causes")

    # ---- Data
    head = ["View", "Branch", "Period", "Key", "Source", "Column", "Old", "New", "Difference", "Difference %",
            "Status", "Cause", "Cause detail"]
    data.append(head)
    style_header(data, 1, len(head))
    keys_old, keys_new = defaultdict(set), defaultdict(set)
    for i, r in enumerate(rows, start=2):
        branch, period = split_key(r["view"], r["key"])
        code = cause(r["view"], r["column"], r["source"], r["status"], branch, period)
        old = float(r["old"]) if r["old"] != "" else None
        new = float(r["new"]) if r["new"] != "" else None
        data.append([r["view"], int(branch) if branch.isdigit() else branch, period, r["key"], r["source"],
                     r["column"], old, new,
                     f'=IF(OR(G{i}="",H{i}=""),"",H{i}-G{i})',
                     f'=IF(OR(G{i}="",H{i}=""),"",IF(G{i}=0,IF(H{i}=0,0,""),(H{i}-G{i})/ABS(G{i})))',
                     r["status"], code, f'=IFERROR(VLOOKUP(L{i},Causes!$A:$B,2,FALSE),"")'])
        if r["status"] != "only new":
            keys_old[r["view"]].add(r["key"])
        if r["status"] != "only old":
            keys_new[r["view"]].add(r["key"])
    last = len(rows) + 1
    for row in data.iter_rows(min_row=2, max_row=last):
        for cell in row:
            cell.font = Font(name=FONT, size=9)
        row[6].number_format = row[7].number_format = row[8].number_format = NUM
        row[9].number_format = PCT
        if row[10].value != "equal":
            row[10].fill = DIFF_FILL
    for col, width in zip("ABCDEFGHIJKLM", (26, 8, 11, 48, 9, 34, 16, 16, 16, 12, 10, 12, 60)):
        data.column_dimensions[col].width = width
    data.freeze_panes = "A2"
    data.auto_filter.ref = f"A1:M{last}"

    # ---- Causes
    legend.append(["Cause", "Explanation"])
    style_header(legend, 1, 2)
    for code, text in CAUSES.items():
        legend.append([code, text])
    for row in legend.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name=FONT, size=10)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    legend.column_dimensions["A"].width = 14
    legend.column_dimensions["B"].width = 120

    # ---- Summary
    views = list(dict.fromkeys(r["view"] for r in rows))
    summary["A1"] = "bsc views: old server (172.22.25.165) against new (172.22.25.214)"
    summary["A1"].font = Font(name=FONT, bold=True, size=13)
    summary["A2"] = ("Compared: the months 2026-08 and 2026-09 and the year 2025 (scripts/compare_bsc.ps1). Values are "
                     "summed per key; a cell differs above 0.5% relative (1 absolute below 1). Counts are formulas over "
                     "the Data sheet; key counts come from the export.")
    summary["A2"].font = Font(name=FONT, italic=True, size=9)
    sh = ["View", "Keys old", "Keys new", "Cells compared", "Equal", "Differs", "Only old", "Only new",
          "Differs: oasis rows", "Differs: fusion rows", "Differs: both rows", "% equal"]
    summary.append([])
    summary.append(sh)
    style_header(summary, 4, len(sh))
    for n, v in enumerate(views, start=5):
        q = f'Data!$A:$A,$A{n}'
        summary.append([
            v, len(keys_old[v]), len(keys_new[v]),
            f"=COUNTIFS({q})",
            f'=COUNTIFS({q},Data!$K:$K,"equal")',
            f'=COUNTIFS({q},Data!$K:$K,"differs")',
            f'=COUNTIFS({q},Data!$K:$K,"only old")',
            f'=COUNTIFS({q},Data!$K:$K,"only new")',
            f'=COUNTIFS({q},Data!$K:$K,"differs",Data!$E:$E,"oasis")',
            f'=COUNTIFS({q},Data!$K:$K,"differs",Data!$E:$E,"fusion")',
            f'=COUNTIFS({q},Data!$K:$K,"differs",Data!$E:$E,"both")',
            f"=IF(D{n}=0,\"\",E{n}/D{n})",
        ])
    end = 4 + len(views)
    for row in summary.iter_rows(min_row=5, max_row=end):
        for cell in row:
            cell.font = Font(name=FONT, size=10)
        for cell in row[1:11]:
            cell.number_format = INT
        row[11].number_format = PCT
    summary.cell(row=4, column=2).comment = Comment("Distinct keys present on the old server (from the export, not a formula).", "compare_bsc")
    summary.column_dimensions["A"].width = 30
    for col in "BCDEFGHIJKL":
        summary.column_dimensions[col].width = 13
    summary.freeze_panes = "A5"

    # ---- Headline KPIs
    measures = [
        ("Net revenue 2025 (SAR)", "vw_financial", "net_revenue_actual", "2025*", "SAME"),
        ("Operating expenses 2025 (SAR)", "vw_financial", "operating_expenses", "2025*", "F6"),
        ("Consumables 2025 (SAR)", "vw_financial", "consumables_actual", "2025*", "SAME"),
        ("Assets Sep-2026 (SAR)", "vw_financial", "assets", "2026-09*", "BS-SEP26"),
        ("Receivables Sep-2026 (SAR)", "vw_financial", "account_receivable_actual", "2026-09*", "BS-SEP26"),
        ("GL revenue movement 2025 (SAR, credit -)", "vw_gl_transactions", "Transaction_Balance", "2025*", "SAME", "*|IS|revenue|*"),
        ("Revenue Aug-2026 (SAR)", "vw_customer", "revenue", "2026-08*", "REV"),
        ("Census Aug-2026", "vw_customer", "census", "2026-08*", "VISITS"),
        ("Admissions Aug-2026", "vw_customer", "admissions", "2026-08*", "ADM"),
        ("Available slots Aug-2026", "vw_customer", "available_slots", "2026-08*", "SLOTS"),
        ("Stays in house Aug-2026", "vw_hospital_beds_utilization", "census", "2026-08*", "BEDS"),
        ("Bed nights Aug-2026", "vw_hospital_beds_utilization", "total_bed_nights", "2026-08*", "BEDS"),
        ("Claims submitted Aug-2026 (SAR)", "vw_rcm", "total_submitted_claims", "2026-08*", "RCM-CLAIMS"),
        ("Claims approved Aug-2026 (SAR)", "vw_rcm", "approved_amount", "2026-08*", "RCM-CLAIMS"),
        ("Pre-auth services Aug-2026", "vw_rcm", "total_orders", "2026-08*", "RCM-PREAUTH"),
        ("Headcount Sep-2026", "vw_work_force", "headcount", "2026-09*", "WF-2026"),
        ("Part-time doctors Sep-2026", "vw_work_force", "pt_doctors_headcount", "2026-09*", "WF-2026"),
        ("Resignations 2025", "vw_work_force", "resignations", "2025*", "SAME"),
    ]
    headline["A1"] = "Headline totals: old against new (SUMIFS over the Data sheet)"
    headline["A1"].font = Font(name=FONT, bold=True, size=13)
    hh = ["Measure", "View", "Column", "Period", "Key filter", "Old", "New", "Difference", "Difference %", "Cause", "Cause detail"]
    headline.append([])
    headline.append(hh)
    style_header(headline, 3, len(hh))
    for n, m in enumerate(measures, start=4):
        label, view, column, period, code = m[:5]
        keyfilter = m[5] if len(m) > 5 else "*"
        crit = f'Data!$A:$A,$B{n},Data!$F:$F,$C{n},Data!$C:$C,$D{n},Data!$D:$D,$E{n}'
        headline.append([label, view, column, period, keyfilter,
                         f"=SUMIFS(Data!$G:$G,{crit})", f"=SUMIFS(Data!$H:$H,{crit})",
                         f"=G{n}-F{n}", f'=IF(F{n}=0,"",H{n}/ABS(F{n}))', code,
                         f'=IFERROR(VLOOKUP(J{n},Causes!$A:$B,2,FALSE),"")'])
    for row in headline.iter_rows(min_row=4, max_row=3 + len(measures)):
        for cell in row:
            cell.font = Font(name=FONT, size=10)
        row[5].number_format = row[6].number_format = row[7].number_format = INT
        row[8].number_format = PCT
        row[10].alignment = Alignment(wrap_text=True, vertical="top")
    headline.cell(row=3, column=5).comment = Comment("Wildcard on the Data key; * = every key. The GL revenue row keeps income-statement revenue accounts only.", "bsc_comparison_xlsx")
    for col, width in zip("ABCDEFGHIJK", (38, 26, 28, 10, 16, 18, 18, 16, 12, 12, 90)):
        headline.column_dimensions[col].width = width
    headline.freeze_panes = "A4"

    # no LibreOffice on HNHANALYTICSSRV to pre-calculate: Excel calculates every formula when the file is opened
    wb.calculation.fullCalcOnLoad = True
    wb.save(out)
    print(f"{out}: {len(rows)} data rows, {len(views)} views")


if __name__ == "__main__":
    main()
